"""
Gera `utils/fluxo_previdenciario.py` a partir da planilha de fluxos atuariais.

Uso (a partir de src/):
    python -m utils.atualizar_fluxo_previdenciario "caminho/Fluxo_Previdênciario.xlsx"

Rodar apenas quando a avaliação atuarial for atualizada.

Linhas usadas por plano:
- "FLUXOS DE SEGURIDADE para o cálculo da duration (Fi)" (RAS), sinal invertido
  -> FLUXO_PREVIDENCIAL_LIQUIDO (benefícios − contribuições de assistidos).
- "RECEBIMENTOS de Ativos **" -> CONTRIBUICOES_ATIVOS (opcional no painel).
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import openpyxl

DESTINO = Path(__file__).resolve().parent / "fluxo_previdenciario.py"


def ler_planilha(caminho: str) -> tuple[dict, dict]:
    ws = openpyxl.load_workbook(caminho, data_only=True).active
    cabecalho = [c.value for c in ws[1]]
    anos = {i: int(v) for i, v in enumerate(cabecalho) if isinstance(v, (int, float))}

    def serie(linha, sinal):
        return {
            ano: round(sinal * float(linha[i]), 2)
            for i, ano in anos.items()
            if isinstance(linha[i], (int, float)) and linha[i] != 0
        }

    liquido, ativos = {}, {}
    for linha in ws.iter_rows(min_row=2, values_only=True):
        plano = str(linha[0]).strip() if linha[0] else None
        descricao = str(linha[1] or "").strip().upper()
        if not plano:
            continue
        if linha[2] == "RAS" and (s := serie(linha, -1)):
            liquido[plano] = s
        elif descricao.startswith("RECEBIMENTOS DE ATIVOS") and (s := serie(linha, 1)):
            ativos[plano] = s
    return liquido, ativos


def _bloco(nome: str, dados: dict) -> list[str]:
    linhas = [f"{nome}: dict[str, dict[int, float]] = {{"]
    for plano, serie in dados.items():
        linhas.append(f"    {plano!r}: {{")
        itens = [f"{ano}: {valor!r}" for ano, valor in sorted(serie.items())]
        for i in range(0, len(itens), 4):
            linhas.append("        " + ", ".join(itens[i:i + 4]) + ",")
        linhas.append("    },")
    linhas.append("}")
    return linhas


def escrever_modulo(liquido: dict, ativos: dict, origem: str) -> None:
    linhas = [
        '"""',
        "Fluxos atuariais anuais por plano (valores fixos).",
        "",
        "ARQUIVO GERADO por utils/atualizar_fluxo_previdenciario.py — não editar à mão.",
        f"Origem: {origem}",
        f"Gerado em: {date.today():%d/%m/%Y}",
        "",
        "FLUXO_PREVIDENCIAL_LIQUIDO: linha RAS com sinal invertido",
        "    (pagamentos de benefícios − contribuições de assistidos).",
        "CONTRIBUICOES_ATIVOS: linha 'RECEBIMENTOS de Ativos' (participantes + patrocinadora).",
        "Chaves = nomes de TESOURARIA da query de recebimentos (de-para 1:1).",
        "Base monetária (real/nominal) e data de referência: ver PREMISSAS em utils/fluxo_caixa.py.",
        '"""',
        "",
        *_bloco("FLUXO_PREVIDENCIAL_LIQUIDO", liquido),
        "",
        *_bloco("CONTRIBUICOES_ATIVOS", ativos),
    ]
    DESTINO.write_text("\n".join(linhas) + "\n", encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Informe o caminho da planilha.")
    liq, atv = ler_planilha(sys.argv[1])
    escrever_modulo(liq, atv, Path(sys.argv[1]).name)
    print(f"{len(liq)} planos com fluxo líquido e {len(atv)} com contribuições de ativos gravados em {DESTINO}")